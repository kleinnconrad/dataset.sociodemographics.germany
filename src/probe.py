"""
Checks candidate sources with the fetch rules and describes what they publish.

The probe is the tool for the rule that every new source is checked before it
is added (CLAUDE.md, sections 8 and 9): for every URL it reports whether
robots.txt allows the request for the pipeline's user agent, the crawl delay,
text-and-data-mining reservations (TDMRep file, header, meta tag and plain
statements in robots.txt), the response (status, redirects, type, size, last
modification) and a description of the content: the title, text and links of
an HTML page, the first lines of a CSV or text file, the keys of a JSON
document, or the members of an archive (xlsx, zip). With a text filter, the
passages of a page's text that match it are listed, e.g. the terms of reuse
in an imprint. Pages are fetched like every website request of the pipeline;
a URL that robots.txt disallows is not fetched.

The source review runs the probe when the workflow ``source_review.yml`` is
started with URLs in the input ``probe``; the description goes to the job
summary, so that sources can be checked live from GitHub Actions.
"""
import io
import json
import re
import zipfile
from html.parser import HTMLParser
from typing import Optional
from urllib.parse import urljoin

from web import MAX_CRAWL_DELAY_S, BlockedError, FetchError, NotFoundError, Response, RobotsDisallowedError, WebClient

MAX_TEXT_CHARS = 2_000
MAX_LINKS = 120
MAX_RULES = 40
MAX_LINES = 15
MAX_LINE_CHARS = 300
MAX_EXCERPTS = 12
MAX_SHEETS = 40
MAX_SHEET_ROWS = 25
EXCERPT_CONTEXT_CHARS = 250
FILE_LINK_PATTERN = re.compile(r"\.(csv|xlsx?|ods|zip|json|pdf|txt|xml)(\?|#|$)", re.IGNORECASE)


class _PageParser(HTMLParser):
    """Collects the title, the visible text, the links and the TDM meta tag of an HTML page."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title = ""
        self.text: list[str] = []
        self.links: list[tuple[str, str]] = []
        self.tdm_meta: Optional[str] = None
        self.robots_meta: Optional[str] = None
        self._skip = 0
        self._in_title = False
        self._href: Optional[str] = None
        self._link_text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, Optional[str]]]) -> None:
        values = {name.lower(): value or "" for name, value in attrs}
        if tag in ("script", "style", "noscript", "template"):
            self._skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag == "a" and values.get("href"):
            self._href = values["href"]
            self._link_text = []
        elif tag == "meta":
            name = values.get("name", "").lower()
            if name == "tdm-reservation":
                self.tdm_meta = values.get("content", "")
            elif name == "robots":
                self.robots_meta = values.get("content", "")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "noscript", "template") and self._skip:
            self._skip -= 1
        elif tag == "title":
            self._in_title = False
        elif tag == "a" and self._href is not None:
            self.links.append((" ".join(" ".join(self._link_text).split()), self._href))
            self._href = None

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        if self._skip:
            return
        self.text.append(data)
        if self._href is not None:
            self._link_text.append(data)


def text_excerpts(text: str, text_filter: re.Pattern[str]) -> list[str]:
    """Returns the passages of a text around the matches of a pattern, overlapping passages merged.

    Args:
        text: The text.
        text_filter: The pattern.

    Returns:
        list[str]: Up to ``MAX_EXCERPTS`` passages.
    """
    spans: list[list[int]] = []
    for match in text_filter.finditer(text):
        start, end = max(0, match.start() - EXCERPT_CONTEXT_CHARS), match.end() + EXCERPT_CONTEXT_CHARS
        if spans and start <= spans[-1][1]:
            spans[-1][1] = max(spans[-1][1], end)
        else:
            spans.append([start, end])
    return [text[start:end] for start, end in spans[:MAX_EXCERPTS]]


def describe_html(response: Response, link_filter: Optional[re.Pattern[str]],
                  text_filter: Optional[re.Pattern[str]] = None) -> list[str]:
    """Describes an HTML page: title, robots and TDM meta tags, text excerpt, matching passages and links.

    Args:
        response: The page.
        link_filter: Only links whose address or text matches are listed; None lists links to files
            first and then the other links.
        text_filter: Passages of the text that match are listed; None lists none.

    Returns:
        list[str]: Lines of the description.
    """
    parser = _PageParser()
    parser.feed(response.text)
    text = " ".join(" ".join(parser.text).split())
    lines = [f"title: {' '.join(parser.title.split())}"]
    if parser.tdm_meta is not None:
        lines.append(f"meta tdm-reservation: {parser.tdm_meta}")
    if parser.robots_meta is not None:
        lines.append(f"meta robots: {parser.robots_meta}")
    lines.append(f"text ({len(text):,} characters): {text[:MAX_TEXT_CHARS]}")
    if text_filter is not None:
        excerpts = text_excerpts(text, text_filter)
        lines.append(f"passages matching the text filter ({len(excerpts)}):")
        lines.extend(f"- ...{excerpt}..." for excerpt in excerpts)
    links = [(label, urljoin(response.url, href)) for label, href in parser.links
             if not href.startswith(("javascript:", "mailto:", "#"))]
    if link_filter is not None:
        links = [(label, href) for label, href in links if link_filter.search(href) or link_filter.search(label)]
    else:
        links.sort(key=lambda link: not FILE_LINK_PATTERN.search(link[1]))
    unique = list(dict.fromkeys(links))
    lines.append(f"links ({len(unique)}{', filtered' if link_filter is not None else ''}):")
    lines.extend(f"- {label[:100]} -> {href}" for label, href in unique[:MAX_LINKS])
    return lines


def describe_workbook(content: bytes) -> list[str]:
    """Describes an xlsx workbook: the name and the first non-empty rows of every sheet.

    Args:
        content: The workbook.

    Returns:
        list[str]: Lines of the description.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as e:
        return [f"xlsx workbook that cannot be read: {e}"]
    try:
        lines = [f"xlsx workbook with {len(workbook.sheetnames)} sheets: " + ", ".join(workbook.sheetnames)]
        for name in workbook.sheetnames[:MAX_SHEETS]:
            rows: list[str] = []
            for row in workbook[name].iter_rows(values_only=True):
                cells = [" ".join(str(value).split()) for value in row if value is not None and str(value).strip()]
                if cells:
                    rows.append(" | ".join(cells)[:MAX_LINE_CHARS])
                if len(rows) >= MAX_SHEET_ROWS:
                    break
            lines.append(f"sheet {name!r}:")
            lines.extend(f"  {row}" for row in rows)
    finally:
        workbook.close()
    return lines


def describe_content(response: Response, link_filter: Optional[re.Pattern[str]] = None,
                     text_filter: Optional[re.Pattern[str]] = None) -> list[str]:
    """Describes the content of a response according to its type.

    Args:
        response: The response.
        link_filter: Filter of the links of an HTML page.
        text_filter: Filter of the passages of the text of an HTML page, a JSON document or a text file.

    Returns:
        list[str]: Lines of the description.
    """
    lines = _describe(response, link_filter, text_filter)
    is_html = lines and lines[0].startswith("title:")
    if text_filter is not None and not is_html and response.content[:2] != b"PK" and response.content[:4] != b"%PDF":
        excerpts = text_excerpts(response.text, text_filter)
        lines.append(f"passages matching the text filter ({len(excerpts)}):")
        lines.extend(f"- ...{' '.join(excerpt.split())}..." for excerpt in excerpts)
    return lines


def _describe(response: Response, link_filter: Optional[re.Pattern[str]],
              text_filter: Optional[re.Pattern[str]]) -> list[str]:
    """Describes the content of a response according to its type (see ``describe_content``)."""
    content = response.content
    kind = response.content_type
    if content[:4] == b"%PDF":
        return [f"PDF document, version {content[5:8].decode('ascii', errors='replace')}"]
    if content[:2] == b"PK":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = archive.namelist()
        except zipfile.BadZipFile:
            return ["damaged zip archive"]
        if any(name.startswith("xl/worksheets/") for name in names):
            return describe_workbook(content)
        return [f"zip archive with {len(names)} members: " + ", ".join(names[:40])]
    if "html" in kind or content.lstrip()[:15].lower().startswith((b"<!doctype html", b"<html")):
        return describe_html(response, link_filter, text_filter)
    if "json" in kind:
        try:
            data = json.loads(content)
        except ValueError:
            return ["invalid JSON"]
        if isinstance(data, dict):
            return [f"JSON object with keys: {', '.join(list(data)[:40])}", json.dumps(data, ensure_ascii=False)[:MAX_TEXT_CHARS]]
        return [f"JSON list with {len(data) if isinstance(data, list) else 0} items",
                json.dumps(data, ensure_ascii=False)[:MAX_TEXT_CHARS]]
    text = response.text
    lines = text.splitlines()
    return [f"text with {len(lines):,} lines; the first lines:"] + [line[:MAX_LINE_CHARS] for line in lines[:MAX_LINES]]


def probe_url(web: WebClient, url: str, link_filter: Optional[str] = None,
              text_filter: Optional[str] = None) -> list[str]:
    """Checks a URL with the fetch rules and describes the response.

    Args:
        web: The HTTP client.
        url: The URL.
        link_filter: Regular expression for the links of an HTML page to list.
        text_filter: Regular expression for the passages of the text of an HTML page to list.

    Returns:
        list[str]: Lines of the report.
    """
    lines = [f"URL: {url}"]
    try:
        rules = web.robots_rules(url)
        allowed = web.allowed(url)
    except FetchError as e:
        return lines + [f"robots.txt not available ({e.reason}); the URL is not fetched"]
    if rules is None:
        lines.append("robots.txt: none (every path allowed)")
    else:
        lines.append(f"robots.txt: {'allows' if allowed else 'disallows'} the URL; crawl delay "
                     f"{rules.crawl_delay or 'none'}; plain-text TDM statement: {'yes' if rules.tdm_statement else 'no'}")
        if rules.crawl_delay and rules.crawl_delay > MAX_CRAWL_DELAY_S:
            lines.append(f"the crawl delay exceeds the maximum of {MAX_CRAWL_DELAY_S:.0f} s of the pipeline; the host "
                         "is not fetched")
        if rules.rules:
            lines.append(f"robots.txt rules for the agent ({len(rules.rules)}): " + ", ".join(
                f"{'allow' if allow else 'disallow'} {pattern}" for allow, pattern, _regex in rules.rules[:MAX_RULES]))
    reserved = web.tdm_reserved(url)
    lines.append(f"TDMRep file: {'reserves' if reserved else 'no reservation'}")
    if not allowed:
        return lines
    try:
        response = web.get(url)
    except RobotsDisallowedError:
        return lines + ["not fetched: disallowed by robots.txt"]
    except BlockedError as e:
        return lines + [f"blocked: {e.reason} (bot protection or refusal)"]
    except NotFoundError as e:
        return lines + [f"not found: {e.reason}"]
    except FetchError as e:
        return lines + [f"failed: {e.reason}"]
    if web.tdm_reserved(url, response):
        lines.append("TDM reservation in the response header")
    headers = response.headers
    lines.append(f"response: HTTP {response.status}, final URL {response.url}"
                 + (f", redirects {list(response.redirects)}" if response.redirects else "")
                 + f", type {response.content_type or 'none'}, {len(response.content):,} bytes"
                 + (f", last modified {headers['last-modified']}" if headers.get("last-modified") else "")
                 + (f", etag {headers['etag']}" if headers.get("etag") else ""))
    pattern = re.compile(link_filter, re.IGNORECASE) if link_filter else None
    text_pattern = re.compile(text_filter, re.IGNORECASE) if text_filter else None
    return lines + describe_content(response, pattern, text_pattern)
