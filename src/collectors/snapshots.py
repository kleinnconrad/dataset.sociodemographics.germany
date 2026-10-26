"""
Stores what sources without a reliable history publish, as snapshots in the dataset ``source_snapshots``.

Sources such as GENESIS tables, the files of federal agencies and church statistics
overwrite their tables with every update, remove older editions or revise figures
without versioning. For every item of ``config/snapshots.toml`` the collectors fetch
the published content, normalize it to columns and rows of text, and compare its
hash with the previous observation. Only changed content writes a record, so the
records of an item are the history of what the source published, observed at the
time of the run (``observed_at``).

Two collectors share the writer:
- ``genesis`` reads tables of the GENESIS databases (GENESIS-Online of the
  Statistisches Bundesamt, Regionaldatenbank Deutschland, Zensus 2022 database)
  through their documented web service (``data/tablefile``, flat CSV in a zip
  archive). The databases need a free account; the token is read from the
  environment variables of ``GENESIS_DATABASES``. Without a token the request is
  sent without credentials, and a refusal is counted as ``no_access``, not as a
  failure.
- ``files`` reads CSV, xlsx, JSON, HTML and PDF files from websites, with the
  fetch rules of every website request (robots.txt, TDM reservations, bot
  protection). A file whose address changes is found through the regular
  expression ``link`` on its ``page``. Conditional requests (ETag,
  Last-Modified) avoid downloading unchanged files. PDF documents and entries
  with ``archive = true`` are stored as files in ``data/files/<snapshot_id>/``.

Cells are stored as text exactly as the source wrote them, so that the snapshots
keep the notation of the source (e.g. decimal commas, markers of suppressed
values); converting them is part of the analysis.
"""
import csv
import io
import json
import logging
import os
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from email.utils import parsedate_to_datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, ClassVar, Optional, Union
from urllib.parse import urljoin

from collectors.base import Collector, CollectorResult, RunContext
from probe import _PageParser
from providers import CONFIG_PATH as PROVIDERS_PATH
from providers import load_providers
from schema import SnapshotRecord
from snapshots import CONFIG_PATH, SnapshotSource, active_snapshots
from state import CollectorEntry, content_hash
from web import BlockedError, FetchError, NotFoundError, RateLimitError, Response, WebClient

logger = logging.getLogger(__name__)

DATASET = "source_snapshots"
FILES_DIR = "files"
REQUEST_TIMEOUT_S = 300.0
MAX_FAILURES_LISTED = 20
# An item whose source refused a request (HTTP 401 or 403, challenge page, CAPTCHA) is not requested again for
# this many days
BLOCKED_RETRY_DAYS = 30
EXTENSIONS = {"csv": "csv", "xlsx": "xlsx", "json": "json", "html": "html", "pdf": "pdf", "genesis": "zip"}


class GenesisError(Exception):
    """A GENESIS database answered with an error status instead of a table.

    Attributes:
        code (int): Status code of the database.
        content (str): Status text.
    """

    def __init__(self, code: int, content: str) -> None:
        super().__init__(f"GENESIS status {code}: {content}")
        self.code = code
        self.content = content

    @property
    def access_denied(self) -> bool:
        """True if the database refused the request for missing or invalid credentials."""
        text = self.content.lower()
        return self.code == 401 or any(marker in text for marker in (
            "anmeld", "passwort", "benutzer", "login", "zugang", "berechtig", "credentials", "password", "unauthor"))


@dataclass(frozen=True)
class GenesisDatabase:
    """A database of the GENESIS software with a REST web service.

    Attributes:
        provider (str): Provider id.
        base_url (str): Base URL of the web service, ending with a slash.
        token_env (str): Environment variable with the token or user name of the account.
        password_env (str): Environment variable with the password of an account without token.
    """
    provider: str
    base_url: str
    token_env: str
    password_env: str

    def credentials(self) -> dict[str, Union[str, bytes]]:
        """Returns the credential headers, or an empty mapping if no token is set."""
        token = os.getenv(self.token_env, "").strip()
        if not token:
            return {}
        return {"username": header_value(token), "password": header_value(os.getenv(self.password_env, "").strip())}


def header_value(text: str) -> Union[str, bytes]:
    """Prepares a text, e.g. a password, for an HTTP header.

    Header values are ASCII. Other characters (e.g. "§" or umlauts in a password) are sent as
    ISO-8859-1 bytes, the encoding in which Java servers such as the GENESIS web services decode
    header bytes, or as UTF-8 if the text has characters outside ISO-8859-1.

    Args:
        text: The text.

    Returns:
        Union[str, bytes]: The text itself if it is ASCII, else its bytes.
    """
    if text.isascii():
        return text
    try:
        return text.encode("iso-8859-1")
    except UnicodeEncodeError:
        return text.encode("utf-8")


GENESIS_DATABASES: dict[str, GenesisDatabase] = {
    database.provider: database for database in (
        GenesisDatabase("destatis", "https://genesis.destatis.de/genesisWS/rest/2020/", "GENESIS_TOKEN",
                        "GENESIS_PASSWORD"),
        GenesisDatabase("regionalstatistik", "https://www.regionalstatistik.de/genesisws/rest/2020/",
                        "REGIONALSTATISTIK_TOKEN", "REGIONALSTATISTIK_PASSWORD"),
        GenesisDatabase("zensus2022", "https://ergebnisse.zensus2022.de/api/rest/2020/", "ZENSUS2022_TOKEN",
                        "ZENSUS2022_PASSWORD"),
    )
}


@dataclass
class PublishedContent:
    """The content of an item as fetched and normalized.

    Attributes:
        url (str): Address that was fetched.
        raw (bytes): The published file.
        columns (list[str]): Column names.
        rows (list[list[Optional[str]]]): Rows as text.
        labels (dict[str, str]): Labels of codes in the rows.
        page_url (Optional[str]): Page from which the address was read.
        title (Optional[str]): Title of the item at the source.
        updated (Optional[datetime]): Last change stated by the source.
        validators (dict[str, str]): ETag and Last-Modified for the next conditional request.
    """
    url: str
    raw: bytes
    columns: list[str] = field(default_factory=list)
    rows: list[list[Optional[str]]] = field(default_factory=list)
    labels: dict[str, str] = field(default_factory=dict)
    page_url: Optional[str] = None
    title: Optional[str] = None
    updated: Optional[datetime] = None
    validators: dict[str, str] = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------

def _cell(value: Any) -> Optional[str]:
    """Converts a parsed cell to text; empty cells become None."""
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float):
        # 15 significant digits, as spreadsheet programs show a number: 8.78, not 8.780000000000001
        return str(int(value)) if value.is_integer() else format(value, ".15g")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def decode_text(raw: bytes, encoding: Optional[str] = None) -> str:
    """Decodes a text file: the configured encoding, else UTF-8 (with or without BOM), else Windows-1252.

    Args:
        raw: The file.
        encoding: Configured encoding.

    Returns:
        str: The text without byte order mark.
    """
    if encoding:
        return raw.decode(encoding).lstrip("﻿")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252")


def parse_csv(raw: bytes, options: dict[str, Any]) -> tuple[list[str], list[list[Optional[str]]]]:
    """Reads a CSV file into columns and rows.

    Lines before the header are skipped: ``skip_lines`` of them, and comment lines that start
    with ``#``; with ``header_pattern``, every line before the first line that matches the
    regular expression. The delimiter is ``delimiter`` or, if not configured, the more frequent
    of ``;`` and ``,`` in the header line.

    Args:
        raw: The file.
        options: ``delimiter``, ``encoding``, ``skip_lines``, ``header_pattern``.

    Returns:
        tuple: Column names and rows.

    Raises:
        ValueError: If the file has no header line.
    """
    lines = decode_text(raw, options.get("encoding")).splitlines()
    lines = lines[int(options.get("skip_lines", 0)):]
    lines = [line for line in lines if not line.startswith("#")]
    if options.get("header_pattern"):
        pattern = re.compile(options["header_pattern"])
        lines = lines[next((index for index, line in enumerate(lines) if pattern.search(line)), len(lines)):]
    if not lines:
        raise ValueError("the CSV file has no header line")
    delimiter = options.get("delimiter") or (";" if lines[0].count(";") >= lines[0].count(",") else ",")
    rows = list(csv.reader(lines, delimiter=delimiter))
    columns = [(_cell(name) or f"column_{position + 1}") for position, name in enumerate(rows[0])]
    return columns, [[_cell(value) for value in row] for row in rows[1:] if any(value.strip() for value in row)]


def parse_xlsx(raw: bytes, options: dict[str, Any]) -> tuple[list[str], list[list[Optional[str]]], str]:
    """Reads one sheet of an xlsx workbook into columns and rows.

    The header starts at row ``header_row`` or, with ``header_pattern``, at the first row whose
    non-empty cells, joined with " | ", match the regular expression. It consists of ``header_rows``
    non-empty rows; the names of a column in these rows are joined with " / ", and a name in a
    merged cell is carried to the right in all but the last header row. The non-empty rows after
    the header are the rows of the content.

    Args:
        raw: The workbook.
        options: ``sheet`` (name; default the first sheet), ``sheet_pattern`` (regular expression for the
            name of the first matching sheet), ``header_row`` (from 1; default 1), ``header_pattern``,
            ``header_rows`` (default 1).

    Returns:
        tuple: Column names, rows and the name of the sheet.

    Raises:
        ValueError: If the sheet does not exist or is empty.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        names = workbook.sheetnames
        if options.get("sheet"):
            name = options["sheet"]
        elif options.get("sheet_pattern"):
            name = next((sheet for sheet in names if re.search(options["sheet_pattern"], sheet)), "")
        else:
            name = names[0]
        if name not in names:
            raise ValueError(f"no sheet {name or options.get('sheet_pattern')!r}; sheets: {', '.join(names[:30])}")
        values = [list(row) for row in workbook[name].iter_rows(values_only=True)]
    finally:
        workbook.close()
    width = max((len(row) for row in values), default=0)
    cells = [[_cell(value) for value in row] + [None] * (width - len(row)) for row in values]
    nonempty = [index for index, row in enumerate(cells) if any(cell is not None for cell in row)]

    def names_of(index: int) -> list[Optional[str]]:
        """The cells of a header row with normalized whitespace (header cells often contain line breaks)."""
        return [" ".join(cell.split()) if cell is not None else None for cell in cells[index]]

    if options.get("header_pattern"):
        pattern = re.compile(options["header_pattern"])
        start = next((index for index in nonempty
                      if pattern.search(" | ".join(cell for cell in names_of(index) if cell is not None))), None)
        if start is None:
            raise ValueError(f"no row of sheet {name!r} matches {options['header_pattern']!r}")
    else:
        start = int(options.get("header_row", 1)) - 1
    header = [index for index in nonempty if index >= start][:int(options.get("header_rows", 1))]
    if not header or header[0] != start:
        raise ValueError(f"sheet {name!r} has no header in row {start + 1}")
    header_cells = [names_of(index) for index in header]
    names: list[list[str]] = [[] for _ in range(width)]
    for number, row in enumerate(header_cells):
        carried: Optional[str] = None
        for position, cell in enumerate(row):
            # A merged cell spans the columns that the header rows below it name
            below = any(later[position] is not None for later in header_cells[number + 1:])
            if cell is None and below:
                cell = carried
            carried = cell if cell is not None else carried
            if cell is not None and (not names[position] or names[position][-1] != cell):
                names[position].append(cell)
    rows = [cells[index] for index in nonempty if index > header[-1]]
    # Columns without a name and without a value (formatting beyond the table) are dropped
    kept = [position for position in range(width) if names[position] or any(row[position] is not None for row in rows)]
    columns = [" / ".join(names[position]) or f"column_{position + 1}" for position in kept]
    return columns, [[row[position] for position in kept] for row in rows], name


def parse_json(raw: bytes) -> tuple[list[str], list[list[Optional[str]]]]:
    """Reads a JSON list of objects (or an object with one list of objects) into columns and rows.

    Args:
        raw: The document.

    Returns:
        tuple: Column names (the keys in order of appearance) and rows; nested values are kept as JSON text.

    Raises:
        ValueError: If the document holds no list of objects.
    """
    data = json.loads(raw)
    if isinstance(data, dict):
        lists = [value for value in data.values() if isinstance(value, list)]
        data = lists[0] if len(lists) == 1 else None
    if not isinstance(data, list) or not all(isinstance(item, dict) for item in data):
        raise ValueError("the JSON document holds no list of objects")
    columns: list[str] = []
    for item in data:
        columns.extend(key for key in item if key not in columns)

    def text(value: Any) -> Optional[str]:
        return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else _cell(value)

    return columns, [[text(item.get(column)) for column in columns] for item in data]


class _TableParser(HTMLParser):
    """Collects the rows of the tables of an HTML page."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._depth = 0
        self._row: Optional[list[str]] = None
        self._cell: Optional[list[str]] = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, Optional[str]]]) -> None:
        if tag == "table":
            self._depth += 1
            if self._depth == 1:
                self.tables.append([])
        elif tag == "tr" and self._depth == 1:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag == "table":
            self._depth = max(0, self._depth - 1)
        elif tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join(" ".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._depth == 1:
            if self._row:
                self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def parse_html_table(raw: bytes, options: dict[str, Any]) -> tuple[list[str], list[list[Optional[str]]]]:
    """Reads one table of an HTML page into columns and rows.

    Args:
        raw: The page.
        options: ``table`` (position of the table on the page, from 0; default 0).

    Returns:
        tuple: Column names (the first row) and rows.

    Raises:
        ValueError: If the page has no such table.
    """
    parser = _TableParser()
    parser.feed(decode_text(raw))
    position = int(options.get("table", 0))
    if position >= len(parser.tables) or not parser.tables[position]:
        raise ValueError(f"the page has {len(parser.tables)} tables; table {position} is missing or empty")
    table = parser.tables[position]
    width = max(len(row) for row in table)
    columns = [table[0][index] if index < len(table[0]) and table[0][index] else f"column_{index + 1}"
               for index in range(width)]
    return columns, [[_cell(row[index]) if index < len(row) else None for index in range(width)] for row in table[1:]]


def parse_genesis_flat(text: str) -> tuple[list[str], list[list[Optional[str]]], dict[str, str], Optional[str]]:
    """Reads a flat CSV table of a GENESIS database (format ``ffcsv``).

    The label columns are moved into a code-to-label map, so that every row keeps only codes and
    values: ``time_code``, ``time``, ``<n>_variable_code`` and ``<n>_variable_attribute_code`` of every
    classifying variable, ``value``, ``value_unit``, ``value_variable_code`` and, if present, the
    quality marker ``value_q``.

    Args:
        text: The CSV text (semicolon-separated).

    Returns:
        tuple: Columns, rows, labels (``<variable>`` and ``<variable>:<attribute>`` for the classifying
        variables, ``time:<code>``, ``value:<code>``, ``statistic:<code>``) and the title of the statistic.

    Raises:
        ValueError: If the text is not a flat GENESIS table.
    """
    reader = csv.DictReader(io.StringIO(text.lstrip("﻿")), delimiter=";")
    header = reader.fieldnames or []
    if "value" not in header or "time" not in header:
        raise ValueError(f"not a flat GENESIS table (columns: {', '.join(header[:8])})")
    variables = sorted({int(name.split("_", 1)[0]) for name in header if re.match(r"^\d+_variable_code$", name)})
    columns = ["time_code", "time"]
    for number in variables:
        columns.extend([f"{number}_variable_code", f"{number}_variable_attribute_code"])
    columns.extend(name for name in ("value", "value_unit", "value_variable_code", "value_q") if name in header)
    labels: dict[str, str] = {}
    title: Optional[str] = None
    rows: list[list[Optional[str]]] = []
    for row in reader:
        if title is None and row.get("statistics_label"):
            title = row["statistics_label"]
            labels[f"statistic:{row.get('statistics_code') or ''}"] = title
        if row.get("time_code") and row.get("time_label"):
            labels.setdefault(f"time:{row['time_code']}", row["time_label"])
        for number in variables:
            variable = row.get(f"{number}_variable_code") or ""
            if variable and row.get(f"{number}_variable_label"):
                labels.setdefault(variable, row[f"{number}_variable_label"])
            attribute = row.get(f"{number}_variable_attribute_code") or ""
            label = row.get(f"{number}_variable_attribute_label") or ""
            if label:
                labels.setdefault(f"{variable}:{attribute}", label)
        if row.get("value_variable_code") and row.get("value_variable_label"):
            labels.setdefault(f"value:{row['value_variable_code']}", row["value_variable_label"])
        rows.append([_cell(row.get(column)) for column in columns])
    return columns, rows, labels, title


def genesis_table(response: Response) -> str:
    """Returns the flat CSV of a GENESIS ``data/tablefile`` response.

    Args:
        response: The response: a zip archive with the CSV, the CSV itself, or a JSON status.

    Returns:
        str: The CSV text.

    Raises:
        GenesisError: If the database answered with an error status.
        ValueError: If the response is neither a table nor a status.
    """
    content = response.content
    if content[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = [name for name in archive.namelist() if name.lower().endswith(".csv")]
            if not names:
                raise ValueError("the zip archive of the table holds no CSV file")
            return decode_text(archive.read(names[0]))
    text = decode_text(content)
    if text.lstrip().startswith("{"):
        try:
            data = json.loads(text)
        except ValueError as e:
            raise ValueError("invalid JSON response of the database") from e
        status = data.get("Status") or {}
        raise GenesisError(int(status.get("Code", -1)), str(status.get("Content", "no status")))
    return text


# ---------------------------------------------------------------------------
# Fetching
# ---------------------------------------------------------------------------

def _post_genesis(web: WebClient, url: str, form: dict[str, str],
                  headers: Optional[dict[str, Union[str, bytes]]]) -> Response:
    """Sends a request to a GENESIS web service; HTTP 401 (no or invalid account) becomes a GenesisError."""
    try:
        return web.post(url, form, headers=headers, timeout_s=REQUEST_TIMEOUT_S)
    except BlockedError as e:
        if e.status == 401:
            raise GenesisError(401, "authentication required (HTTP 401)") from e
        raise


def fetch_genesis(web: WebClient, source: SnapshotSource) -> PublishedContent:
    """Fetches a table of a GENESIS database.

    Args:
        web: The HTTP client.
        source: The snapshot entry (``format = "genesis"``).

    Returns:
        PublishedContent: The table.

    Raises:
        GenesisError: If the database refuses the request or reports an error.
        FetchError: If the request fails.
        ValueError: If the response is not a table.
    """
    database = GENESIS_DATABASES[source.provider]
    headers = database.credentials() or None
    metadata = _post_genesis(web, database.base_url + "metadata/table",
                             {"name": source.table or "", "area": "all", "language": "de"}, headers)
    table_title = genesis_title(metadata)
    if source.expect and not re.search(source.expect, table_title, re.IGNORECASE):
        raise ValueError(f"the title of table {source.table} is {table_title!r}, which does not match {source.expect!r}")
    form = {"name": source.table or "", "area": "all", "compress": "false", "format": "ffcsv", "language": "de",
            "quality": "off", "job": "false", "startyear": str(source.options.get("startyear", "")),
            "endyear": "", "regionalvariable": str(source.options.get("regionalvariable", "")),
            "regionalkey": str(source.options.get("regionalkey", "")), "timeslices": "", "stand": ""}
    url = database.base_url + "data/tablefile"
    response = _post_genesis(web, url, form, headers)
    columns, rows, labels, _statistic = parse_genesis_flat(genesis_table(response))
    return PublishedContent(url=f"{database.base_url}data/tablefile?name={source.table}", raw=response.content,
                            columns=columns, rows=rows, labels=labels, title=table_title,
                            updated=_last_modified(response))


def genesis_title(response: Response) -> str:
    """Reads the title of a table from a GENESIS ``metadata/table`` response.

    Args:
        response: The JSON response.

    Returns:
        str: The title (``Object.Content``) with normalized whitespace.

    Raises:
        GenesisError: If the database answered with an error status.
        ValueError: If the response is not JSON.
    """
    try:
        data = response.json()
    except ValueError as e:
        raise ValueError("the metadata response is not JSON") from e
    status = data.get("Status") or {}
    code = int(status.get("Code", 0) or 0)
    obj = data.get("Object")
    if code not in (0, 22) or not isinstance(obj, dict):
        raise GenesisError(code, str(status.get("Content", "no table metadata")))
    return " ".join(str(obj.get("Content", "")).split())


def genesis_tables(web: WebClient, database: GenesisDatabase, selection: str) -> list[tuple[str, str]]:
    """Lists the tables of a GENESIS database whose code matches a selection (``catalogue/tables``).

    Args:
        web: The HTTP client.
        database: The database.
        selection: Code pattern, e.g. ``12612*``.

    Returns:
        list[tuple[str, str]]: Code and title of every table.

    Raises:
        GenesisError: If the database refuses the request or reports an error.
        FetchError: If the request fails.
        ValueError: If the response is not JSON.
    """
    response = _post_genesis(web, database.base_url + "catalogue/tables",
                             {"selection": selection, "area": "all", "pagelength": "500", "language": "de"},
                             database.credentials() or None)
    try:
        data = response.json()
    except ValueError as e:
        raise ValueError("the catalogue response is not JSON") from e
    status = data.get("Status") or {}
    code = int(status.get("Code", 0) or 0)
    items = data.get("List")
    if code not in (0, 22) and not isinstance(items, list):
        raise GenesisError(code, str(status.get("Content", "no table list")))
    return [(str(item["Code"]), " ".join(str(item.get("Content", "")).split()))
            for item in items or [] if isinstance(item, dict) and item.get("Code")]


def _last_modified(response: Response) -> Optional[datetime]:
    """Reads the Last-Modified header of a response."""
    value = response.headers.get("last-modified")
    if not value:
        return None
    try:
        parsed = parsedate_to_datetime(value)
    except (TypeError, ValueError):
        return None
    return (parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)).astimezone(timezone.utc)


def _get(web: WebClient, url: str, cache: Optional[dict[str, Response]], headers: Optional[dict[str, str]] = None,
         timeout_s: Optional[float] = None) -> Response:
    """Fetches a URL once per run: a successful response is kept in ``cache`` for the other entries of the URL."""
    if cache is not None and url in cache:
        return cache[url]
    response = web.get(url, headers=headers or None, timeout_s=timeout_s)
    if cache is not None and response.status == 200:
        cache[url] = response
    return response


def find_link(web: WebClient, page: str, pattern: str, cache: Optional[dict[str, Response]] = None) -> str:
    """Finds the address of a file on the page that links it.

    Args:
        web: The HTTP client.
        page: The page.
        pattern: Regular expression that the absolute address of the link must contain.
        cache: Responses fetched in this run, shared by the entries of the same page.

    Returns:
        str: The first matching address.

    Raises:
        FetchError: If the page cannot be fetched.
        ValueError: If no link matches.
    """
    response = _get(web, page, cache)
    parser = _PageParser()
    parser.feed(response.text)
    regex = re.compile(pattern)
    for _label, href in parser.links:
        address = urljoin(response.url, href)
        if regex.search(address):
            return address
    raise ValueError(f"no link matching {pattern!r} on {page}")


def fetch_file(web: WebClient, source: SnapshotSource, previous: dict[str, Any],
               cache: Optional[dict[str, Response]] = None) -> Optional[PublishedContent]:
    """Fetches and parses the file of a snapshot entry.

    Several entries can read different sheets or tables of the same file; with ``cache``, the
    file and the page that links it are downloaded once per run.

    Args:
        web: The HTTP client.
        source: The snapshot entry.
        previous: Data of the previous observation (``url``, ``etag``, ``last_modified``).
        cache: Responses fetched in this run.

    Returns:
        Optional[PublishedContent]: The content, or None if the server reports the file as unchanged.

    Raises:
        FetchError: If a request fails or robots.txt disallows it.
        ValueError: If the file cannot be parsed.
    """
    url = source.url or find_link(web, source.page or "", source.link or "", cache)
    headers: dict[str, str] = {}
    if previous.get("url") == url:
        if previous.get("etag"):
            headers["If-None-Match"] = previous["etag"]
        if previous.get("last_modified"):
            headers["If-Modified-Since"] = previous["last_modified"]
    response = _get(web, url, cache, headers, REQUEST_TIMEOUT_S)
    if response.status == 304:
        return None
    if web.tdm_reserved(url, response):
        raise FetchError(url, "text and data mining reserved by the rights holder")
    content = PublishedContent(url=url, raw=response.content, page_url=source.page,
                               title=url.rsplit("/", 1)[-1].split("?")[0], updated=_last_modified(response))
    content.validators = {name: response.headers[header] for name, header in
                          (("etag", "etag"), ("last_modified", "last-modified")) if response.headers.get(header)}
    options = source.options.get(source.format, {})
    if source.format == "csv":
        content.columns, content.rows = parse_csv(response.content, options)
    elif source.format == "xlsx":
        content.columns, content.rows, sheet = parse_xlsx(response.content, options)
        content.title = f"{content.title} ({sheet})"
    elif source.format == "json":
        content.columns, content.rows = parse_json(response.content)
    elif source.format == "html":
        content.columns, content.rows = parse_html_table(response.content, options)
    elif response.content[:4] != b"%PDF":
        raise ValueError("the file is not a PDF document")
    header = f"{content.title}; " + "; ".join(content.columns)
    if source.expect and not re.search(source.expect, header, re.IGNORECASE):
        raise ValueError(f"the header of {url} does not match {source.expect!r}: {header[:200]!r}")
    return content


def snapshot_hash(source: SnapshotSource, content: PublishedContent) -> str:
    """Returns the hash that identifies the published content of an item.

    Args:
        source: The snapshot entry.
        content: The fetched content.

    Returns:
        str: Hash of the columns and rows; of the file itself for a PDF document.
    """
    if source.format == "pdf":
        return content_hash({"file": content.raw.hex()})
    return content_hash({"columns": content.columns, "rows": content.rows})


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------

class SnapshotWriter:
    """Compares fetched content with the previous observation and builds the snapshot records."""

    def __init__(self, ctx: RunContext, licenses: dict[str, str]) -> None:
        """Initializes the writer.

        Args:
            ctx: The run context.
            licenses: License text per provider id.
        """
        self._ctx = ctx
        self._licenses = licenses

    def previous(self, source: SnapshotSource) -> dict[str, Any]:
        """Returns the data stored with the previous observation of an item."""
        entry = self._ctx.seen(DATASET).get(source.id)
        return dict(entry.data) if entry else {}

    def touch(self, source: SnapshotSource) -> None:
        """Records an observation of unchanged content (e.g. HTTP 304)."""
        seen = self._ctx.seen(DATASET)
        entry = seen.get(source.id)
        if entry is not None:
            seen.mark(source.id, entry.hash, self._ctx.today)

    def build(self, source: SnapshotSource, content: PublishedContent) -> Optional[SnapshotRecord]:
        """Builds the record of new or changed content and stores the hash.

        Args:
            source: The snapshot entry.
            content: The fetched content.

        Returns:
            Optional[SnapshotRecord]: The record, or None if the content did not change.
        """
        seen = self._ctx.seen(DATASET)
        item_hash = snapshot_hash(source, content)
        status = seen.classify(source.id, item_hash)
        previous = seen.get(source.id)
        changed = previous.data.get("changed") if previous and status == "unchanged" else self._ctx.today.isoformat()
        data = {"url": content.url, "changed": changed, **content.validators}
        if status == "unchanged":
            seen.mark(source.id, item_hash, self._ctx.today, data=data)
            return None
        archived = self._archive(source, content, item_hash) if source.archive else None
        revision = seen.mark(source.id, item_hash, self._ctx.today, data=data)
        return SnapshotRecord(
            **self._ctx.envelope(DATASET, f"snapshot:{source.id}", revision=revision, source_id=source.provider,
                                 source_url=content.url, license=self._licenses.get(source.provider, "unknown"),
                                 backfilled=False),
            snapshot_id=source.id, name=source.name, topic=source.topic, provider=source.provider, item=source.item(content.url),
            format=source.format, geo_scheme=source.geo_scheme, page_url=content.page_url,
            source_title=content.title, source_updated=content.updated, content_hash=item_hash,
            previous_hash=previous.hash if previous else None, columns=content.columns, rows=content.rows,
            row_count=len(content.rows), labels=content.labels, archived_file=archived, byte_size=len(content.raw),
        )

    def _archive(self, source: SnapshotSource, content: PublishedContent, item_hash: str) -> str:
        """Stores the published file and returns its path relative to the output folder."""
        relative = f"{FILES_DIR}/{source.id}/{self._ctx.today.isoformat()}_{item_hash}.{EXTENSIONS[source.format]}"
        if not self._ctx.dry_run:
            path = self._ctx.settings.output_dir / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content.raw)
        return relative


# ---------------------------------------------------------------------------
# Collectors
# ---------------------------------------------------------------------------

class SnapshotCollector(Collector):
    """Fetches the snapshot entries of some formats (set by the subclasses) and writes changed content."""
    datasets = (DATASET,)
    formats: ClassVar[frozenset[str]]

    def __init__(self, config_path: Path = CONFIG_PATH, providers_path: Path = PROVIDERS_PATH) -> None:
        """Initializes the collector.

        Args:
            config_path: The snapshot configuration.
            providers_path: The provider configuration.
        """
        self._config_path = config_path
        self._providers_path = providers_path

    def fetch(self, ctx: RunContext, source: SnapshotSource, previous: dict[str, Any]) -> Optional[PublishedContent]:
        """Fetches the content of an entry; None if the source reports it as unchanged."""
        raise NotImplementedError

    def collect(self, ctx: RunContext, entry: CollectorEntry) -> CollectorResult:
        """Fetches every due entry and emits a record for new and changed content.

        An entry whose source refuses the request (bot protection) is not requested again for
        ``BLOCKED_RETRY_DAYS`` days.

        Args:
            ctx: The run context.
            entry: Run history; its cursor holds the day of the last check and of the last refusal of every entry.

        Returns:
            CollectorResult: Cursor and counters; an error if every attempted entry failed.
        """
        providers = load_providers(self._providers_path)
        licenses = {provider.id: provider.license for provider in providers}
        sources = active_snapshots(set(self.formats), self._config_path, set(licenses))
        writer = SnapshotWriter(ctx, licenses)
        checked: dict[str, str] = dict(entry.cursor.get("checked") or {})
        blocked: dict[str, str] = dict(entry.cursor.get("blocked") or {})
        counters: dict[str, Any] = {"entries": len(sources), "not_due": 0, "new": 0, "changed": 0, "unchanged": 0,
                                    "no_access": 0, "blocked": 0, "failed": 0}
        failures: list[str] = []
        no_access: list[str] = []
        complete = True
        for position, source in enumerate(sources):
            if ctx.past_deadline():
                counters["deferred"] = len(sources) - position
                complete = False
                break
            last = checked.get(source.id)
            if (source.check_days and last and not ctx.backfill
                    and (ctx.today - date.fromisoformat(last)).days < source.check_days):
                counters["not_due"] += 1
                continue
            refused = blocked.get(source.id)
            if refused and (ctx.today - date.fromisoformat(refused)).days < BLOCKED_RETRY_DAYS:
                counters["blocked"] += 1
                continue
            try:
                content = self.fetch(ctx, source, writer.previous(source))
            except GenesisError as e:
                if e.access_denied:
                    counters["no_access"] += 1
                    no_access.append(source.id)
                else:
                    counters["failed"] += 1
                    failures.append(f"{source.id}: {e}"[:300])
                continue
            except RateLimitError as e:
                logger.warning("%s: rate limit; the remaining entries follow in the next run: %s", source.id, e)
                counters["deferred"] = len(sources) - position
                complete = False
                break
            except (BlockedError, NotFoundError, FetchError, ValueError, zipfile.BadZipFile) as e:
                logger.warning("Snapshot %s failed: %s", source.id, e)
                if isinstance(e, BlockedError):
                    blocked[source.id] = ctx.today.isoformat()
                counters["failed"] += 1
                failures.append(f"{source.id}: {e}"[:300])
                continue
            checked[source.id] = ctx.today.isoformat()
            blocked.pop(source.id, None)
            if content is None:
                writer.touch(source)
                counters["unchanged"] += 1
                continue
            record = writer.build(source, content)
            if record is None:
                counters["unchanged"] += 1
                continue
            counters["new" if record.revision == 0 else "changed"] += 1
            ctx.emit(DATASET, [record])
        ctx.save_seen()
        if failures:
            counters["failed_entries"] = failures[:MAX_FAILURES_LISTED]
        if no_access:
            counters["no_access_entries"] = no_access
        attempted = (counters["entries"] - counters["not_due"] - counters.get("deferred", 0) - counters["no_access"]
                     - counters["blocked"])
        error = f"all {attempted} entries failed" if attempted and counters["failed"] == attempted else None
        logger.info("%s: %d new and %d changed snapshots, %d unchanged, %d without access, %d failed.", self.name,
                    counters["new"], counters["changed"], counters["unchanged"], counters["no_access"],
                    counters["failed"])
        known = {source.id for source in sources}
        cursor = {"checked": {key: value for key, value in checked.items() if key in known},
                  "blocked": {key: value for key, value in blocked.items() if key in known}}
        return CollectorResult(cursor=cursor, counters=counters, complete=complete, error=error)


class GenesisCollector(SnapshotCollector):
    """Snapshots of tables of the GENESIS databases (Destatis, Regionaldatenbank, Zensus 2022)."""
    name = "genesis"
    formats = frozenset({"genesis"})

    def fetch(self, ctx: RunContext, source: SnapshotSource, previous: dict[str, Any]) -> Optional[PublishedContent]:
        """Fetches a GENESIS table."""
        return fetch_genesis(ctx.web, source)


class FilesCollector(SnapshotCollector):
    """Snapshots of files and pages of websites (CSV, xlsx, JSON, HTML tables, PDF)."""
    name = "files"
    formats = frozenset({"csv", "xlsx", "json", "html", "pdf"})

    def __init__(self, config_path: Path = CONFIG_PATH, providers_path: Path = PROVIDERS_PATH) -> None:
        """Initializes the collector.

        Args:
            config_path: The snapshot configuration.
            providers_path: The provider configuration.
        """
        super().__init__(config_path, providers_path)
        self._cache: dict[str, Response] = {}

    def collect(self, ctx: RunContext, entry: CollectorEntry) -> CollectorResult:
        """Fetches every due entry, downloading a file that several entries read once."""
        self._cache = {}
        try:
            return super().collect(ctx, entry)
        finally:
            self._cache = {}

    def fetch(self, ctx: RunContext, source: SnapshotSource, previous: dict[str, Any]) -> Optional[PublishedContent]:
        """Fetches a file, with a conditional request if its address did not change."""
        return fetch_file(ctx.web, source, previous, self._cache)
