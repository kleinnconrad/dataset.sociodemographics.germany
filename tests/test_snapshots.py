"""Tests of the snapshot sources: configuration, parsers and the genesis and files collectors."""
import io
import zipfile
from datetime import date, timedelta

import pytest
from openpyxl import Workbook

from collectors.snapshots import (
    FilesCollector,
    GenesisCollector,
    GenesisError,
    genesis_table,
    parse_csv,
    parse_genesis_flat,
    parse_html_table,
    parse_json,
    parse_xlsx,
)
from conftest import PROVIDERS_FIXTURE, RUN_STARTED, read_dataset
from providers import load_providers
from schema import SnapshotRecord
from snapshots import load_snapshot_config, parse_snapshot
from state import CollectorEntry
from web import Response

FIXTURE = PROVIDERS_FIXTURE.parent / "snapshots.toml"
PROVIDERS = {provider.id for provider in load_providers(PROVIDERS_FIXTURE)}
GENESIS_BASE = "https://ergebnisse.zensus2022.de/api/rest/2020/"
FLAT = ("﻿statistics_code;statistics_label;time_code;time_label;time;1_variable_code;1_variable_label;"
        "1_variable_attribute_code;1_variable_attribute_label;2_variable_code;2_variable_label;"
        "2_variable_attribute_code;2_variable_attribute_label;value;value_unit;value_variable_code;value_variable_label\n"
        "1000A;Bevölkerung kompakt;STAG;Stichtag;2022-05-15;KREISE;Kreise;09162;München, Landeshauptstadt;"
        "REL;Religion;REL01;römisch-katholisch;450000;Anzahl;PRS001;Personen\n"
        "1000A;Bevölkerung kompakt;STAG;Stichtag;2022-05-15;KREISE;Kreise;09162;München, Landeshauptstadt;"
        "REL;Religion;REL02;evangelisch;160000;Anzahl;PRS001;Personen\n")
METADATA = {"Status": {"Code": 0, "Content": "erfolgreich"},
            "Object": {"Code": "1000A-1018", "Content": "Personen:\nReligion"}}


def zipped(name: str, text: str) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(name, text.encode("utf-8"))
    return buffer.getvalue()


def test_configuration_loads_snapshots_and_exclusions() -> None:
    sources, excluded = load_snapshot_config(FIXTURE, PROVIDERS)
    assert [source.id for source in sources] == ["zensus2022.1000A-1018", "bundeswahlleiterin.btw2025_strukturdaten",
                                                 "ekd.kirchenmitglieder_bericht"]
    assert sources[0].item() == "1000A-1018" and sources[0].address == "zensus2022 table 1000A-1018"
    assert sources[1].item("https://wahl.example/dam/btw2025_strukturdaten.csv?__blob=publicationFile") == \
        "btw2025_strukturdaten.csv"
    assert sources[1].options == {"csv": {"skip_lines": 2}}
    assert sources[2].archive and sources[2].check_days == 7
    assert excluded[0].reason == "answers automated requests with HTTP 401"


BASE = {"id": "zensus2022.1000A-1018", "name": "Religion", "topic": "religion", "provider": "zensus2022",
        "format": "genesis", "table": "1000A-1018", "expect": "Religion", "geo_scheme": "ags",
        "added": date(2026, 10, 26)}


@pytest.mark.parametrize("change, message", [
    ({"provider": "unknown"}, "has no entry"),
    ({"id": "other.1000A-1018"}, "must start with the provider"),
    ({"expect": None}, "needs expect"),
    ({"url": "https://x.example/a.csv"}, "needs a table"),
    ({"format": "csv", "table": None, "page": "https://x.example/"}, "url, or page and link"),
    ({"format": "pdf", "table": None, "url": "https://x.example/a.pdf", "archive": False}, "must be archived"),
    ({"topic": "sports"}, "unknown topic"),
    ({"geo_scheme": "postcode"}, "unknown geo_scheme"),
    ({"colour": "red"}, "unknown keys"),
])
def test_invalid_snapshot_entries_are_rejected(change: dict, message: str) -> None:
    raw = {key: value for key, value in {**BASE, **change}.items() if value is not None}
    with pytest.raises(ValueError, match=message):
        parse_snapshot(raw, PROVIDERS)


def test_parse_genesis_flat_moves_labels_into_a_map() -> None:
    columns, rows, labels, title = parse_genesis_flat(FLAT)
    assert columns == ["time_code", "time", "1_variable_code", "1_variable_attribute_code", "2_variable_code",
                       "2_variable_attribute_code", "value", "value_unit", "value_variable_code"]
    assert rows[0] == ["STAG", "2022-05-15", "KREISE", "09162", "REL", "REL01", "450000", "Anzahl", "PRS001"]
    assert labels["KREISE:09162"] == "München, Landeshauptstadt" and labels["REL:REL02"] == "evangelisch"
    assert labels["value:PRS001"] == "Personen" and title == "Bevölkerung kompakt"


def test_genesis_table_reads_zip_archives_and_status_messages() -> None:
    archive = Response(url="u", status=200, headers={}, content=zipped("t_flat.csv", FLAT))
    assert genesis_table(archive).startswith("statistics_code")
    status = Response(url="u", status=200, headers={},
                      content=b'{"Status": {"Code": 15, "Content": "Sie sind nicht berechtigt"}}')
    with pytest.raises(GenesisError) as error:
        genesis_table(status)
    assert error.value.access_denied


def test_parse_csv_skips_header_lines_and_comments() -> None:
    raw = "Titel\n(c) Quelle\n# Kommentar\nWkr-Nr.;Name;Bevölkerung\n1;Flensburg – Schleswig;250,3\n\n".encode("utf-8")
    columns, rows = parse_csv(raw, {"skip_lines": 2})
    assert columns == ["Wkr-Nr.", "Name", "Bevölkerung"]
    assert rows == [["1", "Flensburg – Schleswig", "250,3"]]
    assert parse_csv("a,b\n1,2\n".encode("cp1252"), {})[1] == [["1", "2"]]


def test_parse_xlsx_reads_a_sheet_by_pattern() -> None:
    workbook = Workbook()
    workbook.active.title = "Inhalt"
    sheet = workbook.create_sheet("csv-12613-03")
    sheet.append(["Land", "Woche", "Sterbefälle"])
    sheet.append(["Bayern", 1, 2800.0])
    sheet.append([None, None, None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    columns, rows, name = parse_xlsx(buffer.getvalue(), {"sheet_pattern": "^csv-"})
    assert (columns, rows, name) == (["Land", "Woche", "Sterbefälle"], [["Bayern", "1", "2800"]], "csv-12613-03")
    with pytest.raises(ValueError, match="no sheet"):
        parse_xlsx(buffer.getvalue(), {"sheet": "missing"})


def test_parse_xlsx_joins_a_header_of_several_rows() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Religion"
    sheet.append(["Bevölkerung nach Religionszugehörigkeit"])
    sheet.append([])
    sheet.append(["Amtlicher\nGemeinde-\nschlüssel (AGS)", "Name", "Religion"])
    sheet.append([None, None, "Einwohnerzahl", "Römisch-katholisch", None])
    sheet.append([None, None, "absolut", "absolut", "Anteil in %"])
    sheet.append(["010010000000", "Flensburg, Stadt", 95015, 5213, "5,5%", None, None])
    sheet["F3"] = "Sonstige"
    buffer = io.BytesIO()
    workbook.save(buffer)
    columns, rows, _name = parse_xlsx(buffer.getvalue(), {"header_pattern": "^Amtlicher Gemeinde", "header_rows": 3})
    assert columns == ["Amtlicher Gemeinde- schlüssel (AGS)", "Name", "Religion / Einwohnerzahl / absolut",
                       "Religion / Römisch-katholisch / absolut", "Religion / Römisch-katholisch / Anteil in %",
                       "Sonstige"]
    assert rows == [["010010000000", "Flensburg, Stadt", "95015", "5213", "5,5%", None]]
    assert parse_xlsx(buffer.getvalue(), {"header_row": 3})[0][1:3] == ["Name", "Religion"]
    with pytest.raises(ValueError, match="no row of sheet"):
        parse_xlsx(buffer.getvalue(), {"header_pattern": "^Kreis"})
    with pytest.raises(ValueError, match="no header in row 2"):
        parse_xlsx(buffer.getvalue(), {"header_row": 2})


def test_parse_json_and_html_tables() -> None:
    assert parse_json(b'{"result": [{"a": 1, "b": {"c": 2}}, {"a": 3, "d": null}]}') == (
        ["a", "b", "d"], [["1", '{"c": 2}', None], ["3", None, None]])
    assert parse_json(b'[{"share": 8.780000000000001, "count": 12.0}]') == (["share", "count"], [["8.78", "12"]])
    html = (b"<table><tr><th>Jahr</th><th>Austritte</th></tr><tr><td>2025</td><td>4 512</td></tr></table>"
            b"<table><tr><td>x</td></tr></table>")
    assert parse_html_table(html, {}) == (["Jahr", "Austritte"], [["2025", "4 512"]])
    with pytest.raises(ValueError, match="2 tables"):
        parse_html_table(html, {"table": 3})


def genesis_collector() -> GenesisCollector:
    return GenesisCollector(FIXTURE, PROVIDERS_FIXTURE)


def files_collector() -> FilesCollector:
    return FilesCollector(FIXTURE, PROVIDERS_FIXTURE)


def test_genesis_collector_writes_snapshots_only_for_changed_tables(fake_web, make_ctx, tmp_path) -> None:
    fake_web.add(GENESIS_BASE + "metadata/table", METADATA)
    fake_web.add(GENESIS_BASE + "data/tablefile", zipped("1000A-1018_flat.csv", FLAT), content_type="application/zip")
    ctx = make_ctx()
    first = genesis_collector().collect(ctx, CollectorEntry(first_run=ctx.today))
    assert first.counters["new"] == 1
    later = make_ctx(started=RUN_STARTED + timedelta(days=3))
    assert genesis_collector().collect(later, CollectorEntry(first_run=ctx.today, cursor=first.cursor)).counters[
        "unchanged"] == 1
    revised = FLAT.replace(";450000;", ";450100;")
    fake_web.add(GENESIS_BASE + "data/tablefile", zipped("1000A-1018_flat.csv", revised), content_type="application/zip")
    changed = genesis_collector().collect(make_ctx(started=RUN_STARTED + timedelta(days=7)),
                                          CollectorEntry(first_run=ctx.today, cursor=first.cursor))
    assert changed.counters["changed"] == 1
    records = read_dataset(tmp_path / "data", "source_snapshots")
    assert [record["revision"] for record in records] == [0, 1]
    SnapshotRecord.model_validate(records[1])
    assert records[1]["previous_hash"] == records[0]["content_hash"]
    assert records[1]["source_title"] == "Personen: Religion"
    assert records[1]["rows"][0][6] == "450100" and records[1]["license"].startswith("Datenlizenz Deutschland")


def test_genesis_table_with_another_title_is_not_stored(fake_web, make_ctx) -> None:
    fake_web.add(GENESIS_BASE + "metadata/table", {**METADATA, "Object": {"Content": "Personen: Alter"}})
    ctx = make_ctx()
    result = genesis_collector().collect(ctx, CollectorEntry(first_run=ctx.today))
    assert result.counters["failed"] == 1 and "does not match" in result.counters["failed_entries"][0]
    assert GENESIS_BASE + "data/tablefile" not in fake_web.requested


def test_genesis_http_401_without_token_counts_as_no_access(fake_web, make_ctx, monkeypatch) -> None:
    monkeypatch.delenv("ZENSUS2022_TOKEN", raising=False)
    fake_web.add(GENESIS_BASE + "metadata/table", "Unauthorized", status=401, content_type="text/plain")
    ctx = make_ctx()
    result = genesis_collector().collect(ctx, CollectorEntry(first_run=ctx.today))
    assert result.counters["no_access"] == 1 and result.counters["failed"] == 0
    assert result.cursor["blocked"] == {}


def test_genesis_refusal_without_token_counts_as_no_access(fake_web, make_ctx, monkeypatch) -> None:
    monkeypatch.delenv("ZENSUS2022_TOKEN", raising=False)
    fake_web.add(GENESIS_BASE + "metadata/table", {"Status": {"Code": 15, "Content": "Sie sind nicht berechtigt."}})
    ctx = make_ctx()
    result = genesis_collector().collect(ctx, CollectorEntry(first_run=ctx.today))
    assert result.counters["no_access"] == 1 and result.counters["no_access_entries"] == ["zensus2022.1000A-1018"]
    assert result.error is None


CSV_TEXT = "Strukturdaten\n© Die Bundeswahlleiterin\nWkr-Nr.;Name;Bevölkerung\n1;Flensburg – Schleswig;250,3\n"


def test_files_collector_finds_the_link_and_uses_conditional_requests(fake_web, make_ctx, tmp_path) -> None:
    fake_web.add("https://wahl.example/strukturdaten.html",
                 '<a href="/dam/jcr/abc/btw2025_strukturdaten.csv">CSV</a>')
    file_url = "https://wahl.example/dam/jcr/abc/btw2025_strukturdaten.csv"
    fake_web.add(file_url, CSV_TEXT, content_type="text/csv", headers={"etag": '"v1"'})
    fake_web.add("https://church.example/bericht.pdf", b"%PDF-1.7 report", content_type="application/pdf")
    ctx = make_ctx()
    first = files_collector().collect(ctx, CollectorEntry(first_run=ctx.today))
    assert first.counters["new"] == 2
    records = {record["snapshot_id"]: record for record in read_dataset(tmp_path / "data", "source_snapshots")}
    csv_record = records["bundeswahlleiterin.btw2025_strukturdaten"]
    assert csv_record["columns"] == ["Wkr-Nr.", "Name", "Bevölkerung"] and csv_record["source_url"] == file_url
    assert csv_record["item"] == "btw2025_strukturdaten.csv"
    assert csv_record["page_url"] == "https://wahl.example/strukturdaten.html"
    pdf_record = records["ekd.kirchenmitglieder_bericht"]
    assert pdf_record["rows"] == [] and pdf_record["archived_file"].startswith("files/ekd.kirchenmitglieder_bericht/")
    assert (tmp_path / "data" / pdf_record["archived_file"]).read_bytes() == b"%PDF-1.7 report"

    fake_web.add(file_url, b"", status=304, content_type="text/csv")
    later = make_ctx(started=RUN_STARTED + timedelta(days=3))
    second = files_collector().collect(later, CollectorEntry(first_run=ctx.today, cursor=first.cursor))
    assert second.counters["unchanged"] == 1 and second.counters["not_due"] == 1
    assert len(read_dataset(tmp_path / "data", "source_snapshots")) == 2


def test_files_collector_respects_robots_txt(fake_web, make_ctx) -> None:
    fake_web.add("https://wahl.example/robots.txt", "User-agent: *\nDisallow: /\n", content_type="text/plain")
    fake_web.add("https://church.example/bericht.pdf", b"%PDF-1.7 report", content_type="application/pdf")
    ctx = make_ctx()
    result = files_collector().collect(ctx, CollectorEntry(first_run=ctx.today))
    assert result.counters["failed"] == 1 and "robots.txt" in result.counters["failed_entries"][0]
    assert "https://wahl.example/strukturdaten.html" not in fake_web.requested


def test_files_of_several_entries_are_downloaded_once(fake_web, make_ctx, tmp_path) -> None:
    sources = FIXTURE.read_text(encoding="utf-8").replace(
        'csv.skip_lines = 2', 'csv.skip_lines = 2\nexpect = "Bevölkerung"') + r"""
[[snapshot]]
id = "bundeswahlleiterin.btw2025_strukturdaten_header"
name = "Header of the structural data"
topic = "regional_structure"
provider = "bundeswahlleiterin"
format = "csv"
page = "https://wahl.example/strukturdaten.html"
link = 'strukturdaten\.csv$'
csv.header_pattern = "^Strukturdaten"
geo_scheme = "constituency"
added = 2026-10-26
"""
    config = tmp_path / "snapshots.toml"
    config.write_text(sources, encoding="utf-8")
    fake_web.add("https://wahl.example/strukturdaten.html", '<a href="/dam/btw2025_strukturdaten.csv">CSV</a>')
    fake_web.add("https://wahl.example/dam/btw2025_strukturdaten.csv", CSV_TEXT, content_type="text/csv")
    ctx = make_ctx()
    result = FilesCollector(config, PROVIDERS_FIXTURE).collect(ctx, CollectorEntry(first_run=ctx.today))
    assert result.counters["new"] == 2
    assert fake_web.requested.count("https://wahl.example/dam/btw2025_strukturdaten.csv") == 1
    assert fake_web.requested.count("https://wahl.example/strukturdaten.html") == 1
    records = {record["snapshot_id"]: record for record in read_dataset(tmp_path / "data", "source_snapshots")}
    assert records["bundeswahlleiterin.btw2025_strukturdaten_header"]["columns"] == ["Strukturdaten"]


def test_items_of_a_refusing_source_are_not_requested_again_for_a_while(fake_web, make_ctx) -> None:
    fake_web.add("https://church.example/bericht.pdf", b"denied", status=403, content_type="text/html")
    fake_web.add("https://wahl.example/strukturdaten.html", '<a href="/dam/btw2025_strukturdaten.csv">CSV</a>')
    fake_web.add("https://wahl.example/dam/btw2025_strukturdaten.csv", CSV_TEXT, content_type="text/csv")
    ctx = make_ctx()
    first = files_collector().collect(ctx, CollectorEntry(first_run=ctx.today))
    assert first.counters["failed"] == 1 and first.cursor["blocked"] == {"ekd.kirchenmitglieder_bericht": "2026-10-27"}
    later = make_ctx(started=RUN_STARTED + timedelta(days=10))
    second = files_collector().collect(later, CollectorEntry(first_run=ctx.today, cursor=first.cursor))
    assert second.counters["blocked"] == 1 and second.counters["failed"] == 0
    assert fake_web.requested.count("https://church.example/bericht.pdf") == 1
    fake_web.add("https://church.example/bericht.pdf", b"%PDF-1.7 report", content_type="application/pdf")
    retry = files_collector().collect(make_ctx(started=RUN_STARTED + timedelta(days=31)),
                                      CollectorEntry(first_run=ctx.today, cursor=first.cursor))
    assert retry.counters["new"] == 1 and retry.cursor["blocked"] == {}
